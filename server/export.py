"""CSV / 엑셀(.xlsx) 내보내기.

- CSV: UTF-8 BOM(utf-8-sig)이라 엑셀에서 바로 열어도 한글이 깨지지 않는다.
- 엑셀: 시트별 머리글 서식, 열 너비, 틀 고정을 적용한다.
- 시트 목록에서 primary=True인 시트(없으면 첫 시트)만 CSV로 내보낸다.
- 문자열이 =, +, -, @ 로 시작하면 엑셀이 수식으로 실행할 수 있어 앞에 '를 붙인다(CSV 인젝션 방지).
"""

import csv
import io
from datetime import date
from urllib.parse import quote

from fastapi import HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_TYPE = "text/csv; charset=utf-8"
FORMATS = ("csv", "xlsx")


def check_format(fmt):

    if fmt not in FORMATS:
        raise HTTPException(422, "format은 csv 또는 xlsx여야 합니다.")

    return fmt


def _safe(value):

    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value

    return value


def _sheet(name, headers, rows, primary=False):

    return {"name": name, "headers": headers, "rows": rows, "primary": primary}


def _width(text):
    """한글은 영문의 약 2배 너비로 계산."""

    return sum(2 if ord(ch) > 127 else 1 for ch in str(text))


def build_csv(sheets):

    sheet = next((s for s in sheets if s["primary"]), sheets[0])
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(sheet["headers"])

    for row in sheet["rows"]:
        writer.writerow([_safe(v) for v in row])

    return out.getvalue().encode("utf-8-sig")


def build_xlsx(sheets):

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="0B1F3A")
    header_font = Font(bold=True, color="FFFFFF")

    for sheet in sheets:
        ws = wb.create_sheet(title=sheet["name"][:31])
        ws.append(sheet["headers"])

        for row in sheet["rows"]:
            ws.append([_safe(v) for v in row])

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(sheet["headers"], start=1):
            values = [header] + [r[col_idx - 1] for r in sheet["rows"] if col_idx - 1 < len(r)]
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = min(48, max(10, max(_width(v) for v in values) + 3))

            for row_cells in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell = row_cells[0]

                if isinstance(cell.value, bool):
                    continue
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.0"

        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)

    return buf.getvalue()


def respond(name, fmt, sheets):
    """파일 응답. 파일명은 한글 포함 UTF-8로 전달(RFC 5987)."""

    check_format(fmt)
    filename = f"{name}_{date.today():%Y%m%d}.{fmt}"
    body = build_csv(sheets) if fmt == "csv" else build_xlsx(sheets)

    return Response(
        content=body,
        media_type=CSV_TYPE if fmt == "csv" else XLSX_TYPE,
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"export.{fmt}\"; filename*=UTF-8''{quote(filename)}"
            ),
            "Access-Control-Expose-Headers": "Content-Disposition",
            "Cache-Control": "no-store",
        },
    )


# ------------------------------------------------------------
# 데이터셋별 시트 구성
# ------------------------------------------------------------

def orders_sheets(rows):

    return [_sheet("발주 목록", ["발주번호", "자재", "수량", "단위", "발주일", "납기일", "상태"], [
        [r["po_no"], r["name"], r["quantity"], r["unit"], r["order_date"], r["due_date"], r["status"]]
        for r in rows
    ], primary=True)]


def inventory_sheets(rows):

    return [_sheet("재고 현황", [
        "자재", "단위", "현재 재고", "적정 재고", "부족 수량(적정-현재)", "상태",
        "다음 달 예측 수요", "발주 중 수량", "권장 발주량", "소진 예상(일)", "기준일",
    ], [
        [r["name"], r["unit"], r["current_stock"], r["safety_stock"], r["gap"], r["status"],
         r["next_month_demand"], r["pending_order"], r["recommended_order"], r["days_left"], r["updated_at"]]
        for r in rows
    ], primary=True)]


def forecast_sheets(items, start, end, period_sum, last_year):
    """items: forecast_all() 값 목록(자재 순). 월별 상세가 CSV 대상."""

    prev_start, prev_end = last_year(start, end)
    summary, monthly, models = [], [], []

    for d in items:
        pts = [p for p in d["forecast"] if start <= p["month"] <= end]
        demand = period_sum(d["forecast"], start, end)
        prev = period_sum(d["history"], prev_start, prev_end)

        summary.append([
            d["name"], d["unit"], demand,
            period_sum(d["forecast"], start, end, "low"), period_sum(d["forecast"], start, end, "high"),
            prev, round((demand - prev) / prev * 100, 1) if prev else None,
            d["accuracy"], d["model"]["label"],
        ])

        for p in pts:
            monthly.append([d["name"], d["unit"], p["month"], p["value"], p["low"], p["high"], d["model"]["label"]])

        for c in d["candidates"]:
            models.append([
                d["name"], c["label"],
                round(c["mape_a"] * 100, 2) if c["mape_a"] is not None else None,
                round(c["mape_b"] * 100, 2) if c["mape_b"] is not None else None,
                round(c["mape_avg"] * 100, 2), "선택" if c["selected"] else "",
            ])

    return [
        _sheet("요약", [
            "자재", "단위", "예측 수요(기간)", "하한 합(80%)", "상한 합(80%)",
            "전년 동기 실적", "전년 대비(%)", "예측 정확도(%)", "선택 모델",
        ], summary),
        _sheet("월별 상세", ["자재", "단위", "월", "예측", "하한(80%)", "상한(80%)", "선택 모델"], monthly, primary=True),
        _sheet("모델 비교", [
            "자재", "모델", "MAPE 30개월 학습(%)", "MAPE 36개월 학습(%)", "평균 MAPE(%)", "선택",
        ], models),
    ]


def simulation_sheets(result, params, material_name):

    summary = [
        ["자재", material_name],
        ["기간", f"{params['start']} ~ {params['end']}"],
        ["수요 변동률(%)", params["demand_change"]],
        ["납기 지연(일)", params["delay_days"]],
        ["예상 수요", result["demand"]],
        ["예상 발주량", result["order"]],
        ["예상 결품 수량", result["shortage"]],
        ["예상 재고 소진일", result["days_left"]],
        ["소진일 기준 자재", result["days_left_material"]],
    ]
    monthly = [[m["month"], m["demand"], m["order"], m["shortage"]] for m in result["months"]]

    return [
        _sheet("요약", ["항목", "값"], summary),
        _sheet("월별 결과", ["월", "수요", "발주량", "결품"], monthly, primary=True),
    ]


def scenarios_sheets(rows, material_name, start, end):

    body = [
        [r["label"], r["demand_change"], r["delay_days"], r["demand"], r["order"], r["shortage"], r["days_left"]]
        for r in rows
    ]

    return [_sheet("시나리오 비교", [
        "시나리오", "수요 변동(%)", "납기 지연(일)", "예상 수요", "예상 발주량", "예상 결품", "재고 소진일",
    ], body, primary=True), _sheet("조건", ["항목", "값"], [["자재", material_name], ["기간", f"{start} ~ {end}"]])]


def risks_sheets(data):

    rows = []

    for p in data["projects"]:
        link = p.get("link") or {}

        for item in p["items"]:
            lines = {line["label"]: line["text"] for line in item["lines"]}
            rows.append([
                p["project"], p["status"], p["level"], item["kind_label"], item["title"], item["level"],
                lines.get("위험 원인", ""),
                lines.get("기자재 영향") or lines.get("공사 영향", ""),
                lines.get("수급 문제", ""),
                lines.get("프로젝트 영향", ""),
                link.get("name", ""),
            ])

    return [
        _sheet("위험 알림", [
            "프로젝트", "상태", "프로젝트 등급", "구분", "위험 요인", "등급",
            "위험 원인", "기자재·공사 영향", "수급 문제", "프로젝트 영향", "연결 프로젝트",
        ], rows, primary=True),
        _sheet("기준", ["항목", "값"], [["분석 시각", data["generated_at"]]]),
    ]
